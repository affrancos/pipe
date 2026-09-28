"""Reportes en Excel: resumen por caso y libro de control con una fila por caso."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

COLORES = {"ERROR": "F8CBAD", "ALERTA": "FFE699", "OK": "C6EFCE"}
COLUMNAS_CONTROL = [
    "Caso", "Fecha proceso", "Estado", "Producto", "Convenio", "Cédula", "Nombre",
    "Monto", "Tasa", "Tipo tasa", "Plazo (meses)", "Errores", "Alertas", "Enviado a gestor",
]


def _encabezado(ws, columnas: list[str]) -> None:
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)


def _ajustar_columnas(ws) -> None:
    for col in ws.columns:
        ancho = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 10), 90)


def resumen_caso(resultado: dict, ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    datos = resultado["datos_caso"]
    for etiqueta, valor in [
        ("Caso", resultado["caso"]), ("Estado", resultado["estado"]), ("Fecha proceso", resultado["fecha_proceso"]),
        ("Producto", datos.get("producto")), ("Convenio", datos.get("convenio")), ("Cédula", datos.get("cedula")),
        ("Nombre", datos.get("nombre")), ("Monto", datos.get("monto")), ("Tasa", datos.get("tasa")),
        ("Tipo tasa", datos.get("tipo_tasa")), ("Plazo (meses)", datos.get("plazo_meses")),
    ]:
        ws.append([etiqueta, valor])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    _ajustar_columnas(ws)

    ws = wb.create_sheet("Hallazgos")
    _encabezado(ws, ["Nivel", "Regla", "Detalle"])
    for h in resultado["hallazgos"]:
        ws.append([h["nivel"], h["regla"], h["detalle"]])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=COLORES.get(h["nivel"], "FFFFFF"))
    _ajustar_columnas(ws)

    ws = wb.create_sheet("Documentos")
    _encabezado(ws, ["Archivo", "Tipo", "Puntaje", "Usó OCR", "Campos extraídos"])
    for d in resultado["documentos"]:
        campos = "; ".join(f"{k}={v}" for k, v in d["campos"].items())
        ws.append([d["archivo"], d["tipo"], d["puntaje"], "Sí" if d["uso_ocr"] else "No", campos])
    _ajustar_columnas(ws)

    wb.save(ruta)
    return ruta


def registrar_control(resultado: dict, ruta: Path, enviado: bool = False) -> None:
    """Agrega o actualiza la fila del caso en el libro de control."""
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if ruta.exists():
        wb = load_workbook(ruta)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Control"
        _encabezado(ws, COLUMNAS_CONTROL)

    datos = resultado["datos_caso"]
    niveles = [h["nivel"] for h in resultado["hallazgos"]]
    fila = [
        resultado["caso"], resultado["fecha_proceso"], resultado["estado"], datos.get("producto"),
        datos.get("convenio"), datos.get("cedula"), datos.get("nombre"), datos.get("monto"),
        datos.get("tasa"), datos.get("tipo_tasa"), datos.get("plazo_meses"),
        niveles.count("ERROR"), niveles.count("ALERTA"), "Sí" if enviado else "No",
    ]
    for n in range(2, ws.max_row + 1):
        if str(ws.cell(n, 1).value) == str(resultado["caso"]):
            for i, valor in enumerate(fila, start=1):
                ws.cell(n, i, valor)
            break
    else:
        ws.append(fila)
    _ajustar_columnas(ws)
    wb.save(ruta)
